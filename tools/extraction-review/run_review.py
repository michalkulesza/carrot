"""Auditable review of extraction-v2 captures with Gemini social-text selection.

This tool never fetches web pages or transcribes media.  It uses only selected
capture files, including HTML captures as frozen linked-page snapshots.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import re
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from api.services.extraction_v2.contracts import ExtractionInput, SOURCE_PAYLOAD_ADAPTER
from api.services.extraction_v2.extractor import RecipeEvidenceExtractor
from api.services.extraction_v2.language import LinguaLanguageDetector
from api.services.extraction_v2.orchestrator import ExtractionDependencies, ExtractionOrchestrator
from api.services.extraction_v2.sources import LinkedPage

PREVIEW = 4_000
BOOKKEEPING = {"failed-urls.json"}
REVIEW_FIELDS = ("review_verdict", "error_category", "reviewer_notes", "expected_correction")
_SENSITIVE_ERROR_KEYS = re.compile(
    r"(?:api[_-]?key|authorization|cookie|credential|password|secret|token)", re.IGNORECASE
)
_SECRET_VALUE_PATTERNS = (
    re.compile(r"(?i)(bearer\s+)[^\s,;]+"),
    re.compile(r"(?i)(api[_-]?key\s*[=:]\s*)[^\s,;&]+"),
    re.compile(r"\bAIza[0-9A-Za-z_-]{20,}\b"),
)
_MAX_ERROR_DETAIL_DEPTH = 4
_MAX_ERROR_DETAIL_ITEMS = 40
_MAX_ERROR_DETAIL_TEXT = 4_000


def _redact_error_text(value: str) -> str:
    for pattern in _SECRET_VALUE_PATTERNS:
        value = pattern.sub(lambda match: f"{match.group(1)}[REDACTED]" if match.lastindex else "[REDACTED]", value)
    return value[:_MAX_ERROR_DETAIL_TEXT] + ("… [truncated]" if len(value) > _MAX_ERROR_DETAIL_TEXT else "")


def _safe_error_value(value: Any, *, depth: int = 0, key: str = "") -> Any:
    if _SENSITIVE_ERROR_KEYS.search(key):
        return "[REDACTED]"
    if depth >= _MAX_ERROR_DETAIL_DEPTH:
        return "[truncated]"
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, str):
        return _redact_error_text(value)
    if isinstance(value, dict):
        return {
            str(item_key): _safe_error_value(item_value, depth=depth + 1, key=str(item_key))
            for item_key, item_value in list(value.items())[:_MAX_ERROR_DETAIL_ITEMS]
        }
    if isinstance(value, (list, tuple, set)):
        return [_safe_error_value(item, depth=depth + 1) for item in list(value)[:_MAX_ERROR_DETAIL_ITEMS]]
    if hasattr(value, "model_dump"):
        return _safe_error_value(value.model_dump(mode="json"), depth=depth + 1)
    if hasattr(value, "__dict__"):
        return _safe_error_value(
            {item_key: item_value for item_key, item_value in vars(value).items() if not item_key.startswith("_")},
            depth=depth + 1,
        )
    return _redact_error_text(str(value))


def exception_diagnostic(error: BaseException) -> dict[str, Any]:
    """Return useful provider diagnostics without copying credentials into artifacts."""
    message = str(error)
    provider_message = getattr(error, "message", None)
    if not message and provider_message:
        message = str(provider_message)
    diagnostic: dict[str, Any] = {
        "type": type(error).__name__,
        "message": _redact_error_text(message),
    }
    for name in ("status", "status_code", "code"):
        value = getattr(error, name, None)
        if value is not None:
            diagnostic[name] = _safe_error_value(value, key=name)
    for name in ("details", "errors", "response", "body"):
        value = getattr(error, name, None)
        if value is not None:
            diagnostic[name] = _safe_error_value(value, key=name)
    return diagnostic


def literal_cell(value: object) -> str:
    text = str(value or "")
    return "'" + text if text.lstrip("\x00\t\r\n ").startswith(("=", "+", "-", "@")) else text


def text(value: object, artifact: str = "") -> str:
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    if len(value) > PREVIEW:
        value = f"{value[:PREVIEW]}… [truncated; full artifact: {artifact}]"
    return literal_cell(value)


def sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def url_key(url: str) -> str:
    part = urlsplit(url)
    return urlunsplit((part.scheme, part.netloc, part.path, part.query, ""))


def identifier(path: Path) -> str:
    try:
        name = path.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        name = path.resolve().as_posix()
    return "case-" + sha(name.encode())[:16]


def output_dir(root: Path | None) -> Path:
    base = (root or Path(".local/extraction-review")) / datetime.now(UTC).strftime("%Y%m%dT%H%M%S%fZ")
    base.mkdir(parents=True)
    return base


def discover(values: list[str]) -> tuple[list[Path], list[str]]:
    found: list[Path] = []
    skipped: list[str] = []
    for value in values:
        matches = list(Path().glob(value)) if any(ch in value for ch in "*?[") else [Path(value)]
        if not matches:
            found.append(Path(value))
        for path in matches:
            if path.is_dir():
                for item in sorted(path.glob("*.json")):
                    if item.name.casefold() in BOOKKEEPING:
                        skipped.append(f"{item}: bookkeeping file")
                    else:
                        found.append(item)
            else:
                found.append(path)
    result, seen = [], set()
    for path in sorted(found, key=lambda item: str(item).casefold()):
        key = str(path.resolve()) if path.exists() else str(path.absolute())
        if key in seen:
            skipped.append(f"{path}: duplicate")
        else:
            result.append(path)
            seen.add(key)
    return result, skipped


class Pages:
    def __init__(self, pages: dict[str, LinkedPage], calls: list[dict[str, Any]]) -> None:
        self.pages, self.calls = pages, calls

    async def fetch(self, url: str) -> LinkedPage:
        page = self.pages.get(url_key(url))
        if not page:
            self.calls.append({"stage": "linked_page", "action": "fetch", "status": "missing", "source_url": url, "failure_code": "LINKED_SNAPSHOT_MISSING"})
            raise FileNotFoundError("LINKED_SNAPSHOT_MISSING")
        self.calls.append({"stage": "linked_page", "action": "fetch", "status": "ok", "source_url": url, "final_url": page.final_url})
        return page


class Recorder:
    def __init__(self, calls: list[dict[str, Any]]) -> None:
        self.extractor, self.calls = RecipeEvidenceExtractor(), calls

    async def _extract(self, method: str, source: ExtractionInput):
        started = time.monotonic()
        event: dict[str, Any] = {"stage": method, "action": "extract", "evidence_ids": source.evidence_ids, "input": source.model_dump(mode="json")}
        try:
            result = await getattr(self.extractor, f"extract_{method}")(source)
            event.update(status="ok", candidate=result.model_dump(mode="json"))
            return result
        except Exception as error:
            event.update(status="error", exception_type=type(error).__name__, exception_message=str(error))
            raise
        finally:
            event["duration_ms"] = round((time.monotonic() - started) * 1_000)
            self.calls.append(event)

    async def extract_html(self, source: ExtractionInput): return await self._extract("html", source)
    async def extract_text(self, source: ExtractionInput): return await self._extract("text", source)


class HybridRecorder:
    """Record the final hybrid result rather than leaking the deterministic candidate."""

    def __init__(self, extractor: Any, calls: list[dict[str, Any]]) -> None:
        self.extractor, self.calls = extractor, calls

    @property
    def last_diagnostic(self):
        return self.extractor.last_diagnostic

    async def extract_html(self, source: ExtractionInput):
        return await self.extractor.extract_html(source)

    async def extract_text(self, source: ExtractionInput):
        started = time.monotonic()
        try:
            result = await self.extractor.extract_text(source)
            self.calls.append({
                "stage": "hybrid_text", "action": "select", "status": "fallback" if self.extractor.last_diagnostic else "ok",
                "evidence_ids": source.evidence_ids, "input": source.model_dump(mode="json"), "candidate": result.model_dump(mode="json"),
                "detail": self.extractor.last_diagnostic or "hybrid_selector_applied",
                "duration_ms": round((time.monotonic() - started) * 1_000),
            })
            return result
        except Exception as error:
            self.calls.append({"stage": "hybrid_text", "action": "select", "status": "error", "exception_type": type(error).__name__, "duration_ms": round((time.monotonic() - started) * 1_000)})
            raise


class SelectionRecorder:
    """Keep model selection diagnostics in local review artifacts only."""

    def __init__(self, provider: Any, calls: list[dict[str, Any]], model: str, usage: Any) -> None:
        self.provider, self.calls, self.model, self.usage = provider, calls, model, usage

    async def select(self, source):
        started = time.monotonic()
        before = (self.usage.input_tokens, self.usage.output_tokens, self.usage.calls)
        event = {
            "model": self.model,
            "request_preview": json.dumps([{"id": line.id, "text": line.text} for line in source.lines], ensure_ascii=False),
        }
        try:
            selection = await self.provider.select(source)
            event.update(status="ok", response_preview=selection.model_dump_json())
            return selection
        except Exception as error:
            event.update(status="error", error=exception_diagnostic(error))
            raise
        finally:
            event["latency_ms"] = round((time.monotonic() - started) * 1000)
            event["usage"] = {
                "input_tokens": self.usage.input_tokens - before[0],
                "output_tokens": self.usage.output_tokens - before[1],
                "responses_with_usage": self.usage.calls - before[2],
            }
            self.calls.append(event)


def recipe_rows(recipe: dict[str, Any] | None) -> tuple[str, str, str, list[dict[str, Any]]]:
    groups: list[str] = []; ingredients: list[str] = []; steps: list[str] = []; items: list[dict[str, Any]] = []
    for component_number, component in enumerate((recipe or {}).get("components", []), 1):
        group = component.get("name") or "Main"
        if component.get("ingredients"):
            groups.append(group)
        for number, item in enumerate(component.get("ingredients", []), 1):
            ingredients.append(f"{group}: {item['text']}"); items.append({**item, "type": "ingredient", "group": group, "component": component_number, "order": number})
        for number, item in enumerate(component.get("steps", []), 1):
            steps.append(f"{component_number}.{number}. {item['text']}"); items.append({**item, "type": "step", "group": group, "component": component_number, "order": number})
    return "\n".join(dict.fromkeys(groups)), "\n".join(ingredients), "\n".join(steps), items


def review_components(recipe: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Return the extracted component structure without review-only evidence fields."""

    components: list[dict[str, Any]] = []
    for component in (recipe or {}).get("components", []):
        result: dict[str, Any] = {}
        if component.get("name") is not None:
            result["name"] = component["name"]
        if component.get("ingredients"):
            result["ingredients"] = [{"text": item["text"]} for item in component["ingredients"]]
        if component.get("steps"):
            result["steps"] = [item["text"] for item in component["steps"]]
        if result:
            components.append(result)
    return components


def testable_expectation(result: dict[str, Any]) -> dict[str, Any]:
    """Project a review result into the copy-ready regression expectation schema."""

    recipe = result.get("report_recipe")
    if recipe is None:
        return {
            "capture_errors": ", ".join(result.get("capture_errors", [])),
            "detected_languages": result.get("detected_languages", ""),
            "recipe": None,
        }
    nutrition = recipe.get("nutrition") or {}
    return {
        "capture_errors": ", ".join(result.get("capture_errors", [])),
        "detected_languages": result.get("detected_languages", ""),
        "recipe": {
            "title": recipe.get("title"),
            "yield": recipe.get("yield_servings"),
            "total_time_minutes": recipe.get("total_time_minutes"),
            "nutrition": {
                "calories": nutrition.get("calories"),
                "protein": nutrition.get("protein"),
                "fat": nutrition.get("fat"),
                "carbohydrates": nutrition.get("carbohydrates"),
            },
            "components": review_components(recipe),
        },
    }


def display_yield(recipe: dict[str, Any] | None) -> str:
    return str((recipe or {}).get("yield_servings") or (recipe or {}).get("yield_text") or "")


def readable_recipe_item(item: dict[str, Any], wording_key: str) -> dict[str, Any]:
    """Turn spreadsheet-shaped recipe rows back into useful JSON objects."""

    def parsed(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value

    links = parsed(item.get("links", "[]"))
    return {
        wording_key: item["original_wording"],
        "group": item["group"],
        "order": item["order"],
        "evidence_ids": item["evidence_ids"].split(", ") if item.get("evidence_ids") else [],
        "locator": item.get("locator") or None,
        "link_url": item.get("link_url") or None,
        "links": links if isinstance(links, list) else [],
        "references": parsed(item.get("references", "[]")),
    }


async def review(path: Path, pages: dict[str, LinkedPage], run: Path, hybrid: bool = True) -> dict[str, Any]:
    result: dict[str, Any] = {"case_id": identifier(path), "input_file": str(path), "stages": [], "model_calls": []}
    if not path.exists(): return {**result, "review_status": "runner_error", "runner_error": "INPUT_NOT_FOUND"}
    raw = path.read_bytes(); result["input_sha256"] = sha(raw)
    try: envelope = json.loads(raw); payload = SOURCE_PAYLOAD_ADAPTER.validate_python(envelope)
    except Exception as error: return {**result, "review_status": "runner_error", "runner_error": "INVALID_ENVELOPE", "exception_type": type(error).__name__, "exception_message": str(error)}
    artifact = f"cases/{sha(raw)[:16]}-input.json"; (run / artifact).parent.mkdir(parents=True, exist_ok=True); (run / artifact).write_bytes(raw)
    calls: list[dict[str, Any]] = []
    extractor = Recorder(calls)
    if hybrid and payload.kind.value == "social":
        # HTML and explicitly offline reviews never need Gemini configuration.
        from api.services.extraction_v2.gemini_selection import GeminiTextSelectionProvider, HybridTextExtractor
        from api.config import settings
        from api.services.gemini import UsageTracker
        usage = UsageTracker()
        provider = SelectionRecorder(
            GeminiTextSelectionProvider(usage=usage), result["model_calls"], settings.gemini_text_selection_model, usage,
        )
        selected_extractor = HybridRecorder(HybridTextExtractor(extractor, provider), calls)
    else:
        selected_extractor = extractor
    outcome = await ExtractionOrchestrator(ExtractionDependencies(selected_extractor, LinguaLanguageDetector(), Pages(pages, calls), transcription_provider=None)).extract(envelope)
    data = outcome.model_dump(mode="json"); recipe = data.get("recipe")
    candidate_recipe = next((call.get("candidate") for call in reversed(calls) if call.get("candidate")), None)
    report_recipe = recipe or candidate_recipe
    groups, ingredients, instructions, items = recipe_rows(report_recipe)
    limitations = ["LINKED_SNAPSHOT_MISSING"] if any(call.get("failure_code") == "LINKED_SNAPSHOT_MISSING" for call in calls) else []
    limitations.extend(call["detail"] for call in calls if call.get("stage") == "hybrid_text" and call.get("status") == "fallback")
    transcript = getattr(getattr(payload, "audio", None), "transcript", None)
    if transcript and outcome.outcome != "complete": limitations.append("AUDIO_MODEL_DISABLED")
    languages = data.get("evidence", []); trace = data.get("trace", [])
    nutrition = (report_recipe or {}).get("nutrition") or {}
    result.update({
        "source_url": str(payload.source_url), "kind": payload.kind.value,
        "capture_status": payload.capture.status, "capture_errors": payload.capture.errors,
        "review_status": "limited" if limitations else "finished", "review_limitations": limitations,
        "outcome": data["outcome"], "issue_codes": data.get("issue_codes", []),
        "failure_reason": data.get("reason", ""), "failed_stage": data.get("failed_stage", ""),
        "detected_languages": ", ".join(dict.fromkeys((item["language"].get("code") or "undetermined") for item in languages)),
        "language_details": [f"{item['id']}:{item['language'].get('code') or 'undetermined'}:{item['language'].get('confidence')}" for item in languages],
        "completed_at_stage": "complete" if outcome.outcome == "complete" else "not_complete",
        "decision_reason": trace[-1]["event"] if trace else "", "transcript_available": bool(transcript),
        "transcript_used": any(item.get("event") == "audio_merged" for item in trace),
        "transcript_status": "available" if transcript else "absent", "fallback_reason": ", ".join(limitations),
        "title": (report_recipe or {}).get("title", ""), "yield": display_yield(report_recipe),
        "yield_source_text": (report_recipe or {}).get("yield_text", ""),
        "total_time_minutes": (report_recipe or {}).get("total_time_minutes", ""),
        "total_time_source_text": (report_recipe or {}).get("total_time_text", ""), "calories": nutrition.get("calories", ""),
        "protein": nutrition.get("protein", ""), "fat": nutrition.get("fat", ""),
        "carbohydrates": nutrition.get("carbohydrates", ""), "ingredient_groups": groups,
        "ingredients": ingredients, "instructions": instructions, "recipe_items": items,
        "components": review_components(report_recipe),
        "report_recipe": report_recipe,
        "artifact_path": artifact, "duration_ms": 0, "outcome_data": data,
    })
    result["stages"] = [{"sequence": index, **event} for index, event in enumerate([*calls, *trace], 1)]
    (run / f"cases/{sha(raw)[:16]}-result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def notes(path: Path | None) -> dict[tuple[str, str], dict[str, str]]:
    if not path or not path.exists(): return {}
    sheet = load_workbook(path, read_only=True)["Summary"]; rows = sheet.iter_rows(values_only=True); headers = {value: i for i, value in enumerate(next(rows))}
    return {(str(row[headers["case_id"]]), str(row[headers["input_sha256"]])): {field: str(row[headers[field]] or "") for field in REVIEW_FIELDS} for row in rows}


def report(results: list[dict[str, Any]], run: Path, previous: dict[tuple[str, str], dict[str, str]], csv_enabled: bool) -> Path:
    book = Workbook(); book.remove(book.active)
    definitions = {"Summary": ["case_id","input_file","input_sha256","source_url","kind","capture_status","capture_errors","review_status","review_limitations","outcome","issue_codes","failure_reason","failed_stage","detected_languages","language_details","completed_at_stage","decision_reason","transcript_available","transcript_used","transcript_status","fallback_reason","title","yield","yield_source_text","total_time_minutes","total_time_source_text","calories","protein","fat","carbohydrates","ingredient_groups","ingredients","instructions","duration_ms","artifact_path",*REVIEW_FIELDS], "Stages": ["case_id","sequence","stage","action","status","evidence_ids","source_url","input_preview","candidate_preview","detail","failure_code","exception_type","exception_message","duration_ms"], "Evidence": ["case_id","evidence_id","kind","source_url","creator_verified","source_text","locator","language","confidence"], "Recipe items": ["case_id","type","group","component","order","original_wording","evidence_ids","locator","link_url","references","links"], "Model calls": ["case_id","status","model","request_preview","response_preview","latency_ms","usage","error"]}
    sheets = {name: book.create_sheet(name) for name in definitions}
    for name, headers in definitions.items():
        sheet=sheets[name]; sheet.append(headers); sheet.freeze_panes="B2"; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="1F4E78")
    for result in results:
        saved=previous.get((result["case_id"],result.get("input_sha256","")),{})
        sheets["Summary"].append([text(result.get(key,""),result.get("artifact_path","")) for key in definitions["Summary"][:-4]]+[text(saved.get(key,"")) for key in REVIEW_FIELDS])
        for event in result.get("stages",[]): sheets["Stages"].append([text(result["case_id"]),event.get("sequence"),text(event.get("stage","")),text(event.get("action",event.get("event",""))),text(event.get("status","")),text(", ".join(event.get("evidence_ids",[]))),text(event.get("source_url","")),text(json.dumps(event.get("input",""),ensure_ascii=False),result.get("artifact_path","")),text(json.dumps(event.get("candidate",""),ensure_ascii=False)),text(event.get("detail",event.get("event",""))),text(event.get("failure_code","")),text(event.get("exception_type","")),text(event.get("exception_message","")),event.get("duration_ms","")])
        for evidence in result.get("outcome_data",{}).get("evidence",[]):
            language=evidence["language"]; sheets["Evidence"].append([text(result["case_id"]),text(evidence["id"]),text(evidence["kind"]),text(evidence["source_url"]),evidence.get("author_verified",""),text(evidence["text"],result.get("artifact_path","")),text(evidence.get("locator","")),text(language.get("code") or "undetermined"),language.get("confidence","")])
        for item in result.get("recipe_items",[]): sheets["Recipe items"].append([text(result["case_id"]),item["type"],text(item["group"]),item["component"],item["order"],text(item["text"]),text(", ".join(item.get("evidence_ids",[]))),text(item.get("locator","")),text(item.get("link_url","")),text(json.dumps(item.get("references",[]),ensure_ascii=False)),text(json.dumps(item.get("links",[]),ensure_ascii=False))])
        for call in result.get("model_calls", []):
            sheets["Model calls"].append([
                text(result["case_id"]), *[
                    text(json.dumps(call.get(key), ensure_ascii=False) if key == "error" and isinstance(call.get(key), dict) else call.get(key, ""))
                    for key in definitions["Model calls"][1:]
                ],
            ])
    for sheet in sheets.values():
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=30
            for cell in column: cell.alignment=Alignment(wrap_text=True,vertical="top")
    path=run/"review.xlsx"; book.save(path)
    tables = {
        sheet.title: [dict(zip(next(sheet.iter_rows(values_only=True)), row, strict=True)) for row in sheet.iter_rows(min_row=2, values_only=True)]
        for sheet in book.worksheets
    }
    components_by_case = {result["case_id"]: result.get("components", []) for result in results}
    model_calls_by_case = {result["case_id"]: result.get("model_calls", []) for result in results}
    cases = []
    for summary in tables["Summary"]:
        case_id = summary["case_id"]
        items = [item for item in tables["Recipe items"] if item["case_id"] == case_id]
        cases.append({
            "case_id": case_id,
            "summary": summary,
            "ingredients": [
                readable_recipe_item(item, "ingredient") for item in items if item["type"] == "ingredient"
            ],
            "steps": [
                readable_recipe_item(item, "step") for item in items if item["type"] == "step"
            ],
            "components": components_by_case.get(case_id, []),
            "evidence": [item for item in tables["Evidence"] if item["case_id"] == case_id],
            "stages": [item for item in tables["Stages"] if item["case_id"] == case_id],
            # Keep structured provider diagnostics in JSON; the spreadsheet cell is a
            # readable serialized representation for Excel users.
            "model_calls": model_calls_by_case.get(case_id, []),
        })
    (run / "review.json").write_text(json.dumps({"version": 2, "cases": cases}, ensure_ascii=False, indent=2), encoding="utf-8")
    testable = {
        "version": 1,
        "url": results[0].get("source_url", "") if len(results) == 1 else None,
        "expectations": {
            Path(result["input_file"]).name: testable_expectation(result)
            for result in results
            if result.get("input_file")
        },
    }
    (run / "review-testable.json").write_text(json.dumps(testable, ensure_ascii=False, indent=2), encoding="utf-8")
    if csv_enabled:
        for sheet in book.worksheets:
            with (run/f"{sheet.title.lower().replace(' ','-')}.csv").open("w",newline="",encoding="utf-8-sig") as file: csv.writer(file).writerows(sheet.values)
    return path


async def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument("inputs",nargs="*"); parser.add_argument("--manifest",type=Path); parser.add_argument("--output",type=Path); parser.add_argument("--linked-fixtures",type=Path); parser.add_argument("--resume-from",type=Path); parser.add_argument("--csv",action="store_true"); parser.add_argument("--offset",type=int,default=0,help="skip this many deterministically sorted inputs"); parser.add_argument("--limit",type=int,help="process at most this many inputs after --offset"); parser.add_argument("--audio-model",choices=("off","replay","live"),default="off"); parser.add_argument("--text-selector",choices=("off","live"),default="live",help="Gemini social-text selection (default: live); off runs offline"); parser.add_argument("--recordings",type=Path); parser.add_argument("--allow-paid-models",action="store_true",help=argparse.SUPPRESS); parser.add_argument("--max-model-calls",type=int)
    args=parser.parse_args()
    if bool(args.inputs)==bool(args.manifest): parser.error("provide either capture inputs or --manifest")
    if args.manifest: parser.error("legacy manifest mode is no longer supported; use capture envelopes")
    if args.audio_model!="off": parser.error("audio replay/live adapters are not implemented; offline mode is the supported review mode")
    if args.max_model_calls is not None: parser.error("--max-model-calls is not supported for bounded retrying selector calls")
    if args.offset < 0: parser.error("--offset must be zero or greater")
    if args.limit is not None and args.limit <= 0: parser.error("--limit must be positive")
    paths,skips=discover(args.inputs)
    paths=paths[args.offset:] if args.limit is None else paths[args.offset:args.offset + args.limit]
    if not paths: parser.error("no JSON envelopes discovered")
    run=output_dir(args.output); all_paths=[*paths]
    if args.linked_fixtures: all_paths.extend(discover([str(args.linked_fixtures)])[0])
    pages={}
    for path in all_paths:
        try:
            data=json.loads(path.read_text(encoding="utf-8"))
            if data.get("kind")=="html": pages[url_key(data["source_url"])]=LinkedPage(data["source_url"],data["source_url"],data["html"])
        except Exception: pass
    results=[]
    for path in paths:
        try: results.append(await review(path,pages,run,args.text_selector == "live"))
        except Exception as error: results.append({"case_id":identifier(path),"input_file":str(path),"review_status":"runner_error","runner_error":"UNEXPECTED_RUNNER_ERROR","exception_type":type(error).__name__,"exception_message":str(error),"stages":[],"model_calls":[]})
    workbook=report(results,run,notes(args.resume_from),args.csv)
    (run/"run.json").write_text(json.dumps({"version":2,"created_at":datetime.now(UTC).isoformat(),"inputs":[str(path) for path in paths],"discovery_skips":skips,"audio_model":"off","text_selector":args.text_selector,"workbook":str(workbook)},indent=2),encoding="utf-8")
    print(f"Reviewed {len(results)} input(s): {workbook}")
    return 1 if any(item.get("review_status")!="finished" for item in results) else 0


if __name__=="__main__": raise SystemExit(asyncio.run(main()))
