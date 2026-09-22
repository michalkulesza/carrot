"""Hybrid review artifacts describe actual selections and visible fallback."""

import asyncio
import importlib.util
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from api.services.extraction_v2 import gemini_selection
from api.services.extraction_v2.contracts import AudioExtractionInput, ExtractedRecipe, LanguageResult, TextSelection
from api.services.gemini import UsageTracker


def runner():
    spec = importlib.util.spec_from_file_location("hybrid_review_runner", "tools/extraction-review/run_review.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.asyncio
@pytest.mark.parametrize("invalid", [False, True])
async def test_review_exports_actual_selection_model_and_fallback(tmp_path, monkeypatch, invalid):
    module = runner()
    monkeypatch.setattr(module, "LinguaLanguageDetector", lambda: SimpleNamespace(
        detect=lambda _: LanguageResult(code="en", confidence=1),
    ))

    class Selector:
        def __init__(self, *, usage):
            self.usage = usage

        async def select(self, source):
            self.usage.add(SimpleNamespace(usage_metadata=SimpleNamespace(prompt_token_count=5, candidates_token_count=3)))
            return TextSelection.model_validate({
                "components": [{"ingredient_line_ids": ["line:999"]}] if invalid else [],
            })

    monkeypatch.setattr(gemini_selection, "GeminiTextSelectionProvider", Selector)
    capture = {
        "schema_version": 1, "kind": "social", "source_url": "https://www.instagram.com/reel/example/",
        "capture": {"status": "complete", "errors": []},
        "scrapecreators_response": {"data": {"xdt_shortcode_media": {
            "edge_media_to_caption": {"edges": [{"node": {"text": "Ingredients:\n1 onion"}}]},
            "owner": {"username": "cook"},
        }}},
        "comments": [], "audio": {"status": "unavailable"},
    }
    path = tmp_path / "capture.json"
    path.write_text(json.dumps(capture), encoding="utf-8")
    result = await module.review(path, {}, tmp_path)

    event = next(event for event in result["stages"] if event.get("stage") == "hybrid_text")
    assert event["input"]["content"] == "Ingredients:\n1 onion"
    assert event["status"] == ("fallback" if invalid else "ok")
    if invalid:
        assert result["report_recipe"]["components"][0]["ingredients"][0]["text"] == "1 onion"
        assert result["review_status"] == "limited"
        assert "hybrid_selector_invalid_or_failed_fell_back" in result["fallback_reason"]
    else:
        assert result["report_recipe"]["components"] == []
        assert result["text_selection"] == {"components": [], "yield_line_ids": [], "nutrition": {"line_ids": []}}
    assert result["model_calls"][0]["usage"]["input_tokens"] == 5
    workbook = module.report([result], tmp_path, {}, False)
    book = load_workbook(workbook, read_only=True)
    assert book["Model calls"].max_row == 2
    book.close()
    artifact = json.loads((tmp_path / "review.json").read_text(encoding="utf-8"))
    assert artifact["cases"][0]["model_calls"]
    assert artifact["cases"][0]["text_selection"] == result["text_selection"]
    testable = json.loads((tmp_path / "review-testable.json").read_text(encoding="utf-8"))
    exported = testable["expectations"]["capture.json"]
    if invalid:
        assert "text_selection" not in exported
    else:
        assert exported["text_selection"] == result["text_selection"]


@pytest.mark.asyncio
@pytest.mark.parametrize("extra,enabled,audio_enabled", [
    ([], True, True),
    (["--text-selector", "off"], False, True),
    (["--audio-model", "off"], True, False),
])
async def test_cli_defaults_to_gemini_and_saved_transcript_extraction_without_approval_flag(tmp_path, monkeypatch, extra, enabled, audio_enabled):
    module = runner()
    capture = tmp_path / "capture.json"
    capture.write_text("{}", encoding="utf-8")
    seen = []

    async def review(path, pages, run, hybrid, audio):
        seen.append((hybrid, audio))
        return {"review_status": "finished"}

    monkeypatch.setattr(module, "review", review)
    monkeypatch.setattr(module, "report", lambda *args: tmp_path / "review.xlsx")
    monkeypatch.setattr("sys.argv", ["run_review.py", str(capture), "--output", str(tmp_path / "out"), *extra])
    assert await module.main() == 0
    assert seen == [(enabled, audio_enabled)]


@pytest.mark.asyncio
async def test_audio_review_request_times_out_and_records_the_failure(monkeypatch):
    module = runner()

    class BlockingAudioExtractor:
        def __init__(self, **_kwargs):
            pass

        async def extract_audio(self, _source):
            await asyncio.sleep(10)

    from api.services.extraction_v2 import adapters
    monkeypatch.setattr(adapters, "GeminiAudioEvidenceExtractor", BlockingAudioExtractor)
    calls = []
    recorder = module.AudioRecorder(calls, "gemini-test", UsageTracker(), timeout_seconds=0.01)

    with pytest.raises(TimeoutError):
        await recorder.extract_audio(AudioExtractionInput(transcript="Cook it.", retained_recipe=ExtractedRecipe()))

    assert calls[0]["status"] == "error"
    assert calls[0]["error"]["type"] == "TimeoutError"


@pytest.mark.asyncio
async def test_default_html_review_does_not_construct_gemini_provider(tmp_path, monkeypatch):
    module = runner()
    monkeypatch.setattr(module, "LinguaLanguageDetector", lambda: SimpleNamespace(
        detect=lambda _: LanguageResult(code="en", confidence=1),
    ))

    def forbidden(**kwargs):
        raise AssertionError("HTML review must not create a Gemini provider")

    monkeypatch.setattr(gemini_selection, "GeminiTextSelectionProvider", forbidden)
    capture = tmp_path / "html.json"
    capture.write_text(json.dumps({
        "schema_version": 1, "kind": "html", "source_url": "https://example.com/recipe",
        "capture": {"status": "complete", "errors": []},
        "html": "<h2>Ingredients</h2><p>1 onion</p><h2>Instructions</h2><p>Cook.</p>",
    }), encoding="utf-8")
    result = await module.review(capture, {}, tmp_path)
    assert result["model_calls"] == []
    assert result["outcome"] == "complete"
